from pathlib import Path

import fitz
import pytesseract

from PIL import Image
from docx import Document
from pptx import Presentation

from backend.services.embeddings import generate_embeddings
from backend.services.chunker import create_chunks
from backend.services.vector_store import (
    save_document,
    save_chunks_batch,
    update_document_status
)


# ============================================================
# Configuration
# ============================================================

UPLOAD_DIR = Path("data/uploads")

UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True
)


# Windows Tesseract installation
TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


# ============================================================
# Save uploaded file
# ============================================================

def save_uploaded_file(uploaded_file):
    """
    Save an uploaded file to the uploads directory.

    This function is kept for the existing
    Streamlit workflow.
    """

    file_path = UPLOAD_DIR / uploaded_file.name

    with open(file_path, "wb") as file:
        file.write(uploaded_file.getbuffer())

    return file_path


# ============================================================
# Extract PDF
# ============================================================

def extract_pdf(file_path, document_id=None):
    """
    Extract text from a PDF.

    Normal PDFs:
        Uses PyMuPDF text extraction.

    Scanned PDFs:
        Automatically uses Tesseract OCR on pages
        where no text is found.
    """

    document = fitz.open(file_path)

    pages = []

    total_pages = len(document)

    print(
        f"Processing PDF: {total_pages} pages"
    )

    for page_number, page in enumerate(document):

        current_page = page_number + 1

        # ----------------------------------------------------
        # Update extraction progress
        # ----------------------------------------------------

        if document_id is not None:

            update_document_status(
                document_id,
                status="processing",
                processing_stage="extracting",
                current_page=current_page,
                total_pages=total_pages
            )

        # ----------------------------------------------------
        # Try normal PDF text extraction first
        # ----------------------------------------------------

        text = page.get_text().strip()

        # ----------------------------------------------------
        # If no text exists, use OCR
        # ----------------------------------------------------

        if not text:

            print(
                f"OCR required for page "
                f"{current_page}/{total_pages}"
            )

            if document_id is not None:

                update_document_status(
                    document_id,
                    status="processing",
                    processing_stage="ocr",
                    current_page=current_page,
                    total_pages=total_pages
                )

            # Render PDF page as image
            pix = page.get_pixmap(
                matrix=fitz.Matrix(2, 2),
                alpha=False
            )

            # Convert PyMuPDF image to PIL image
            image = Image.frombytes(
                "RGB",
                [pix.width, pix.height],
                pix.samples
            )

            # Run OCR
            text = pytesseract.image_to_string(
                  image,
                 config="--psm 6"
            ).strip()

            # Free memory
            del image
            del pix

        # ----------------------------------------------------
        # Store page
        # ----------------------------------------------------

        pages.append(
            {
                "page": current_page,
                "text": text
            }
        )

    document.close()

    return pages


# ============================================================
# Extract DOCX
# ============================================================

def extract_docx(file_path):
    """
    Extract text from a DOCX file.
    """

    document = Document(file_path)

    text = []

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text.append(
                paragraph.text
            )

    return [
        {
            "page": 1,
            "text": "\n".join(text)
        }
    ]


# ============================================================
# Extract PPTX
# ============================================================

def extract_pptx(file_path):
    """
    Extract text from every slide of a PPTX file.
    """

    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides
    ):

        text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text.append(
                    shape.text
                )

        slides.append(
            {
                "page": slide_number + 1,
                "text": "\n".join(text)
            }
        )

    return slides


# ============================================================
# Process document
# ============================================================

def process_document(
    file_path,
    document_id=None
):
    """
    Extract text from:

    - PDF
    - DOCX
    - PPTX
    - TXT
    """

    file_path = Path(file_path)

    extension = file_path.suffix.lower()

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":

        return extract_pdf(
            file_path,
            document_id=document_id
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    elif extension == ".docx":

        return extract_docx(
            file_path
        )

    # --------------------------------------------------------
    # PPTX
    # --------------------------------------------------------

    elif extension == ".pptx":

        return extract_pptx(
            file_path
        )

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    elif extension == ".txt":

        with open(
            file_path,
            "r",
            encoding="utf-8"
        ) as file:

            return [
                {
                    "page": 1,
                    "text": file.read()
                }
            ]

    # --------------------------------------------------------
    # Unsupported file
    # --------------------------------------------------------

    else:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )


# ============================================================
# Process and store document
# ============================================================

def process_and_store_document(
    file_path,
    pages,
    document_id=None
):
    """
    Process document using:

    1. Text chunking
    2. Batch embeddings
    3. Batch PostgreSQL inserts
    4. Processing status tracking
    """

    file_path = Path(file_path)

    # --------------------------------------------------------
    # Create document record if needed
    # --------------------------------------------------------

    if document_id is None:

        document_id = save_document(
            filename=file_path.name,
            file_type=file_path.suffix.lower()
        )

    try:

        # ----------------------------------------------------
        # Mark document as processing
        # ----------------------------------------------------

        update_document_status(
            document_id,
            status="processing",
            processing_stage="chunking"
        )

        # ----------------------------------------------------
        # Create chunks
        # ----------------------------------------------------

        chunks = create_chunks(
            pages
        )

        total_chunks = len(chunks)

        update_document_status(
            document_id,
            status="processing",
            processing_stage="embedding",
            total_chunks=total_chunks,
            processed_chunks=0,
            error_message=None
        )

        # ----------------------------------------------------
        # No chunks found
        # ----------------------------------------------------

        if not chunks:

            error_message = (
                "No readable text was found in the "
                "document. The document may be empty, "
                "corrupted, or contain unsupported content."
            )

            update_document_status(
                document_id,
                status="failed",
                total_chunks=0,
                processed_chunks=0,
                error_message=error_message
            )

            raise ValueError(
                error_message
            )

        # ----------------------------------------------------
        # Batch configuration
        # ----------------------------------------------------

        embedding_batch_size = 64
        database_batch_size = 250

        processed_chunks = 0

        # ----------------------------------------------------
        # Process chunks in batches
        # ----------------------------------------------------

        for start in range(
            0,
            total_chunks,
            database_batch_size
        ):

            batch_chunks = chunks[
                start:start + database_batch_size
            ]

            # ------------------------------------------------
            # Extract text
            # ------------------------------------------------

            texts = [
                chunk["text"]
                for chunk in batch_chunks
            ]

            # ------------------------------------------------
            # Generate embeddings
            # ------------------------------------------------

            embeddings = generate_embeddings(
                texts,
                batch_size=embedding_batch_size
            )

            # ------------------------------------------------
            # Save chunks
            # ------------------------------------------------

            save_chunks_batch(
                document_id,
                batch_chunks,
                embeddings
            )

            # ------------------------------------------------
            # Update progress
            # ------------------------------------------------

            processed_chunks += len(
                batch_chunks
            )

            update_document_status(
                document_id,
                status="processing",
                processing_stage="embedding",
                total_chunks=total_chunks,
                processed_chunks=processed_chunks,
                error_message=None
            )

            print(
                f"Processed "
                f"{processed_chunks}/"
                f"{total_chunks} chunks"
            )

        # ----------------------------------------------------
        # Mark completed
        # ----------------------------------------------------

        update_document_status(
            document_id,
            status="completed",
            processing_stage="completed",
            total_chunks=total_chunks,
            processed_chunks=total_chunks,
            error_message=None
        )

        print(
            f"Document {document_id} completed "
            f"with {total_chunks} chunks"
        )

        # ----------------------------------------------------
        # Return result
        # ----------------------------------------------------

        return {
            "document_id": document_id,
            "chunks": total_chunks
        }

    except Exception as error:

        # ----------------------------------------------------
        # Mark failed
        # ----------------------------------------------------

        update_document_status(
            document_id,
            status="failed",
            error_message=str(error)
        )

        raise